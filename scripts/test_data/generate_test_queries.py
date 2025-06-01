"""
Script to generate test documents for querying the RAG system.
These documents represent new contracts and policies that users would upload
to compare against historical documents in the RAG system's knowledge base.
"""

import os
from fpdf import FPDF
from docx import Document
import random
from datetime import datetime, timedelta
import uuid

def generate_contract_data():
    """Generate realistic current contract data for test queries."""
    companies = [
        "AcmeCorp", "BetaFinance", "DeltaTech", "EpsilonSolutions", 
        "GammaInnovations", "OmegaSystems", "SigmaEnterprises"
    ]
    
    equipment_types = [
        "Office Equipment",
        "Manufacturing Machinery",
        "IT Infrastructure",
        "Medical Equipment",
        "Construction Equipment",
        "Energy Systems",
        "Logistics Equipment",
        "Food Processing Machinery",
        "Educational Technology",
        "Research Equipment"
    ]
    
    # Generate data with current date
    data = {
        "company": random.choice(companies),
        "date": datetime.now().strftime("%Y-%m-%d"),
        "equipment": random.choice(equipment_types),
        "lease_term": random.choice([24, 36, 48]),
        "monthly_payment": random.randint(1000, 5000),
        "agreement_number": f"LEA-{datetime.now().strftime('%Y%m')}-{random.randint(1, 999):03d}"
    }
    return data

def generate_policy_data():
    """Generate realistic current policy data for test queries."""
    policy_types = [
        "Credit Assessment",
        "Payment Terms",
        "Collection Procedures",
        "Risk Management",
        "Customer Onboarding"
    ]
    
    # Generate data with current date
    data = {
        "policy_type": random.choice(policy_types),
        "effective_date": datetime.now().strftime("%Y-%m-%d"),
        "version": f"v{random.randint(1, 5)}.{random.randint(0, 9)}",
        "policy_number": f"POL-{datetime.now().strftime('%Y%m')}-{random.randint(1, 999):03d}"
    }
    return data

def create_contract_pdf(data):
    """Create a PDF contract with proper formatting."""
    pdf = FPDF()
    pdf.add_page()
    
    # Title
    pdf.set_font('Arial', 'B', 16)
    pdf.cell(0, 10, 'EQUIPMENT LEASING AGREEMENT', ln=True, align='C')
    pdf.ln(10)
    
    # Agreement Details
    pdf.set_font('Arial', '', 12)
    pdf.cell(0, 10, f'Agreement Date: {data["date"]}', ln=True)
    pdf.cell(0, 10, f'Agreement Number: {data["agreement_number"]}', ln=True)
    pdf.ln(5)
    
    # Parties
    pdf.set_font('Arial', 'B', 12)
    pdf.cell(0, 10, 'PARTIES', ln=True)
    pdf.set_font('Arial', '', 12)
    pdf.cell(0, 10, f'Lessor: Financial Services Corp.', ln=True)
    pdf.cell(0, 10, f'Lessee: {data["company"]}', ln=True)
    pdf.ln(5)
    
    # Terms and Conditions
    pdf.set_font('Arial', 'B', 12)
    pdf.cell(0, 10, 'TERMS AND CONDITIONS', ln=True)
    pdf.set_font('Arial', '', 12)
    
    terms = [
        f"1. LEASE TERM: {data['lease_term']} months from the commencement date.",
        f"2. MONTHLY PAYMENT: EUR {data['monthly_payment']:,.2f}, due on the first day of each month.",
        f"3. EQUIPMENT: {data['equipment']} as specified in Schedule A.",
        "4. INSURANCE: Lessee shall maintain comprehensive insurance coverage.",
        "5. MAINTENANCE: Lessee responsible for routine maintenance.",
        "6. EARLY TERMINATION: 3 months' notice required with early termination fee.",
        "7. RENEWAL: Option to renew for additional 12 months at market rate.",
        "8. DEFAULT: Late payment fee of 5% after 5 business days.",
        "9. WARRANTIES: Equipment provided 'as is' with standard manufacturer warranty.",
        "10. CONFIDENTIALITY: Both parties agree to maintain confidentiality."
    ]
    
    for term in terms:
        pdf.multi_cell(0, 10, term)
        pdf.ln(2)
    
    # Save the PDF
    os.makedirs('data/test_documents', exist_ok=True)
    filename = f'data/test_documents/{data["company"]}_new_contract.pdf'
    pdf.output(filename)
    return filename

def create_policy_pdf(data):
    """Create a PDF policy document with proper formatting."""
    pdf = FPDF()
    pdf.add_page()
    
    # Title
    pdf.set_font('Arial', 'B', 16)
    pdf.cell(0, 10, f'{data["policy_type"].upper()} POLICY', ln=True, align='C')
    pdf.ln(10)
    
    # Policy Details
    pdf.set_font('Arial', '', 12)
    pdf.cell(0, 10, f'Effective Date: {data["effective_date"]}', ln=True)
    pdf.cell(0, 10, f'Policy Number: {data["policy_number"]}', ln=True)
    pdf.cell(0, 10, f'Version: {data["version"]}', ln=True)
    pdf.ln(5)
    
    # Policy Content
    pdf.set_font('Arial', 'B', 12)
    pdf.cell(0, 10, 'POLICY CONTENT', ln=True)
    pdf.set_font('Arial', '', 12)
    
    # Generate policy-specific content
    if data["policy_type"] == "Credit Assessment":
        content = [
            "1. CREDIT SCORE REQUIREMENTS",
            "   - Minimum credit score: 650",
            "   - No recent bankruptcies",
            "   - Clean payment history",
            "",
            "2. FINANCIAL ASSESSMENT",
            "   - Minimum annual revenue: EUR 500,000",
            "   - Positive cash flow for last 12 months",
            "   - Acceptable debt-to-income ratio",
            "",
            "3. DOCUMENTATION REQUIREMENTS",
            "   - Financial statements",
            "   - Tax returns",
            "   - Bank statements"
        ]
    elif data["policy_type"] == "Payment Terms":
        content = [
            "1. STANDARD PAYMENT TERMS",
            "   - Net 30 for standard customers",
            "   - Net 15 for preferred customers",
            "   - Net 45 for large enterprise customers",
            "",
            "2. LATE PAYMENT PROCEDURES",
            "   - 5% late fee after 5 days",
            "   - Account review after 30 days",
            "   - Collection process after 60 days",
            "",
            "3. PAYMENT METHODS",
            "   - Bank transfer",
            "   - Credit card (2% fee)",
            "   - Direct debit"
        ]
    else:
        content = [
            "1. POLICY OBJECTIVES",
            "   - Ensure consistent application",
            "   - Maintain compliance",
            "   - Protect company interests",
            "",
            "2. IMPLEMENTATION GUIDELINES",
            "   - Regular review and updates",
            "   - Staff training requirements",
            "   - Documentation standards",
            "",
            "3. COMPLIANCE REQUIREMENTS",
            "   - Regular audits",
            "   - Reporting procedures",
            "   - Record keeping"
        ]
    
    for line in content:
        pdf.multi_cell(0, 10, line)
        pdf.ln(2)
    
    # Save the PDF
    os.makedirs('data/test_documents', exist_ok=True)
    filename = f'data/test_documents/{data["policy_type"].lower().replace(" ", "_")}_new_policy.pdf'
    pdf.output(filename)
    return filename

def generate_test_documents(num_docs=5):
    """Generate a mix of new contracts and policies for testing the RAG system."""
    generated_files = []
    
    print("Generating test documents for RAG system queries...")
    print("These documents represent new contracts and policies that users would upload")
    print("to compare against historical documents in the RAG system.\n")
    
    # Generate 3 contracts and 2 policies
    for i in range(num_docs):
        if i < 3:  # Generate contracts
            data = generate_contract_data()
            filename = create_contract_pdf(data)
        else:  # Generate policies
            data = generate_policy_data()
            filename = create_policy_pdf(data)
            
        generated_files.append(filename)
        print(f"Generated test document: {filename}")
    
    return generated_files

if __name__ == "__main__":
    files = generate_test_documents()
    print("\nGenerated test files:")
    for file in files:
        print(f"- {file}")
    
    print("\nThese documents can be used to test the RAG system by:")
    print("1. Uploading them through the Streamlit interface")
    print("2. Asking questions about their content")
    print("3. Comparing them against historical documents in the RAG system") 